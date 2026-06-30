#!/usr/bin/env python3
import argparse
import csv
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path


RUNTIME_VERSION = "0.1.0"
COMMON_CHINESE_OFFICE_FONTS = [
    "宋体",
    "SimSun",
    "仿宋",
    "仿宋_GB2312",
    "FangSong",
    "楷体",
    "楷体_GB2312",
    "KaiTi",
    "黑体",
    "SimHei",
    "Microsoft YaHei",
    "微软雅黑",
    "方正小标宋简体",
    "Times New Roman",
    "TimesNewRoman",
    "新罗马",
]


def fail(message, code=1):
    print(json.dumps({"ok": False, "error": message}, ensure_ascii=False), file=sys.stderr)
    raise SystemExit(code)


def read_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def write_json(path, data):
    path = Path(path)
    if path.parent:
        path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")


def require_docx():
    try:
        import docx
        return docx
    except Exception as exc:
        fail(f"python-docx is required: {exc}")


def require_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except Exception as exc:
        fail(f"openpyxl is required: {exc}")


def import_status():
    modules = ["docx", "openpyxl", "pandas", "xlsxwriter", "xlrd", "mammoth", "lxml", "PIL"]
    result = {}
    for name in modules:
        try:
            __import__(name)
            result[name] = "ok"
        except Exception as exc:
            result[name] = f"missing: {type(exc).__name__}: {exc}"
    return result


def capabilities(args):
    root = Path(__file__).resolve().parents[1]
    path = root / "capabilities.json"
    data = read_json(path)
    data["runtime_version"] = RUNTIME_VERSION
    data["imports"] = import_status()
    data["binaries"] = {"soffice": shutil.which("soffice") or "", "fc-match": shutil.which("fc-match") or ""}
    data["fonts"] = font_matches()
    print(json.dumps(data, ensure_ascii=False, indent=2))


def font_matches():
    fc_match = shutil.which("fc-match")
    result = {}
    if not fc_match:
        return result
    for family in COMMON_CHINESE_OFFICE_FONTS:
        try:
            proc = subprocess.run(
                [fc_match, "-f", "%{family}|%{file}\n", family],
                text=True,
                capture_output=True,
                timeout=5,
                check=False,
            )
            result[family] = proc.stdout.strip()
        except Exception as exc:
            result[family] = f"error: {type(exc).__name__}: {exc}"
    return result


def list_fonts(args):
    print(json.dumps({"ok": True, "fonts": font_matches()}, ensure_ascii=False, indent=2))


def apply_docx_styles(doc, spec):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    section = doc.sections[0]
    page = spec.get("page", {})
    margin = float(page.get("margin_inches", 0.8))
    section.top_margin = Inches(margin)
    section.bottom_margin = Inches(margin)
    section.left_margin = Inches(margin)
    section.right_margin = Inches(margin)

    normal = doc.styles["Normal"]
    normal.font.name = spec.get("font", "Arial")
    normal.font.size = Pt(float(spec.get("body_pt", 10.5)))
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color in [
        ("Title", 24, "111827"),
        ("Heading 1", 15, "1F4E79"),
        ("Heading 2", 12.5, "374151"),
    ]:
        style = doc.styles[style_name]
        style.font.name = spec.get("font", "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)

    title_align = spec.get("title_align", "left")
    return WD_ALIGN_PARAGRAPH.CENTER if title_align == "center" else WD_ALIGN_PARAGRAPH.LEFT


def add_docx_table(doc, table_spec):
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    headers = table_spec.get("headers", [])
    rows = table_spec.get("rows", [])
    col_count = max(len(headers), *(len(r) for r in rows), 1)
    table = doc.add_table(rows=1 if headers else 0, cols=col_count)
    table.style = table_spec.get("style", "Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    if headers:
        cells = table.rows[0].cells
        for i, value in enumerate(headers):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(255, 255, 255)
            shade_cell(cells[i], table_spec.get("header_fill", "1F4E79"))

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = "" if value is None else str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()


def shade_cell(cell, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def create_docx(args):
    docx = require_docx()
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    spec = read_json(args.spec)
    doc = docx.Document()
    title_alignment = apply_docx_styles(doc, spec)

    title = spec.get("title")
    if title:
        p = doc.add_paragraph(style="Title")
        p.alignment = title_alignment
        p.add_run(str(title))
    subtitle = spec.get("subtitle")
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = title_alignment
        run = p.add_run(str(subtitle))
        run.italic = True

    for section in spec.get("sections", []):
        heading = section.get("heading")
        if heading:
            doc.add_heading(str(heading), level=int(section.get("level", 1)))
        for paragraph in section.get("paragraphs", []):
            doc.add_paragraph(str(paragraph))
        for items in section.get("bullets", []):
            for item in items:
                doc.add_paragraph(str(item), style="List Bullet")
        for table in section.get("tables", []):
            add_docx_table(doc, table)
        if section.get("page_break"):
            doc.add_page_break()

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    inspection = inspect_docx_file(out)
    print(json.dumps({"ok": True, "output": str(out), "inspection": inspection}, ensure_ascii=False))


def inspect_docx_file(path):
    docx = require_docx()
    doc = docx.Document(path)
    paragraphs = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            paragraphs.append({"style": p.style.name if p.style else "", "text": text})
    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        tables.append(rows)
    return {"paragraph_count": len(paragraphs), "table_count": len(tables), "paragraphs": paragraphs[:80], "tables": tables[:20]}


def inspect_docx(args):
    data = inspect_docx_file(args.input)
    if args.out:
        write_json(args.out, data)
    print(json.dumps({"ok": True, "input": args.input, "inspection": data}, ensure_ascii=False))


def edit_docx(args):
    docx = require_docx()
    ops_data = read_json(args.ops)
    ops = ops_data.get("operations", []) if isinstance(ops_data, dict) else ops_data
    doc = docx.Document(args.input)

    for op in ops:
        typ = op.get("type")
        if typ == "replace_text":
            old = str(op.get("old", ""))
            new = str(op.get("new", ""))
            for p in doc.paragraphs:
                if old in p.text:
                    for run in p.runs:
                        run.text = run.text.replace(old, new)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.text = run.text.replace(old, new)
        elif typ == "append_paragraph":
            doc.add_paragraph(str(op.get("text", "")))
        elif typ == "set_table_cell":
            table_idx = int(op.get("table", 0))
            row_idx = int(op.get("row", 0))
            col_idx = int(op.get("col", 0))
            doc.tables[table_idx].rows[row_idx].cells[col_idx].text = str(op.get("value", ""))
        else:
            fail(f"unsupported docx operation: {typ}")

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    inspection = inspect_docx_file(out)
    print(json.dumps({"ok": True, "output": str(out), "inspection": inspection}, ensure_ascii=False))


def normalize_rows(rows):
    return [["" if v is None else v for v in row] for row in rows]


def create_xlsx(args):
    openpyxl = require_openpyxl()
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    spec = read_json(args.spec)
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    thin = Side(style="thin", color="D1D5DB")

    for idx, sheet_spec in enumerate(spec.get("sheets", [])):
        ws = wb.create_sheet(sheet_spec.get("name", f"Sheet{idx + 1}"))
        rows = normalize_rows(sheet_spec.get("rows", []))
        for row in rows:
            ws.append(row)
        for cell_ref, formula in sheet_spec.get("formulas", {}).items():
            ws[cell_ref] = formula
        if rows:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=sheet_spec.get("header_fill", "1F4E79"))
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            ws.freeze_panes = sheet_spec.get("freeze_panes", "A2")
            ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        for col, width in sheet_spec.get("widths", {}).items():
            ws.column_dimensions[col].width = float(width)
        if sheet_spec.get("table") and ws.max_row > 1 and ws.max_column > 1:
            ref = f"A1:{ws.cell(ws.max_row, ws.max_column).coordinate}"
            table = Table(displayName=sheet_spec.get("table_name", f"Table{idx + 1}"), ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            ws.add_table(table)
        for chart_spec in sheet_spec.get("charts", []):
            chart_type = chart_spec.get("type", "bar")
            chart = LineChart() if chart_type == "line" else BarChart()
            chart.title = chart_spec.get("title", "")
            data = Reference(ws, min_col=int(chart_spec["data_min_col"]), min_row=int(chart_spec.get("data_min_row", 1)), max_col=int(chart_spec["data_max_col"]), max_row=int(chart_spec["data_max_row"]))
            cats = Reference(ws, min_col=int(chart_spec["cats_col"]), min_row=int(chart_spec.get("cats_min_row", 2)), max_row=int(chart_spec["data_max_row"]))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.height = float(chart_spec.get("height", 7))
            chart.width = float(chart_spec.get("width", 12))
            ws.add_chart(chart, chart_spec.get("anchor", "G2"))

    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    inspection = inspect_xlsx_file(out)
    print(json.dumps({"ok": True, "output": str(out), "inspection": inspection}, ensure_ascii=False))


def inspect_xlsx_file(path, sample_rows=10):
    openpyxl = require_openpyxl()
    wb = openpyxl.load_workbook(path, data_only=False)
    result = {"sheets": []}
    for ws in wb.worksheets:
        formulas = []
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append({"cell": cell.coordinate, "formula": cell.value})
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, sample_rows), values_only=True):
            rows.append(list(row))
        result["sheets"].append({
            "name": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "headers": rows[0] if rows else [],
            "sample_rows": rows[1:],
            "formulas": formulas[:100],
        })
    return result


def inspect_xlsx(args):
    data = inspect_xlsx_file(args.input)
    if args.out:
        write_json(args.out, data)
    print(json.dumps({"ok": True, "input": args.input, "inspection": data}, ensure_ascii=False))


def edit_xlsx(args):
    openpyxl = require_openpyxl()
    ops_data = read_json(args.ops)
    ops = ops_data.get("operations", []) if isinstance(ops_data, dict) else ops_data
    wb = openpyxl.load_workbook(args.input)
    for op in ops:
        typ = op.get("type")
        if typ == "add_sheet":
            name = op.get("sheet", "Sheet")
            ws = wb.create_sheet(name)
            for row in op.get("rows", []):
                ws.append(row)
        else:
            ws = wb[op.get("sheet", wb.sheetnames[0])]
            if typ in ("set_cell", "set_formula"):
                ws[op["cell"]] = op.get("formula") if typ == "set_formula" else op.get("value")
            elif typ == "append_row":
                ws.append(op.get("values", []))
            elif typ == "set_width":
                ws.column_dimensions[op["column"]].width = float(op.get("width", 12))
            else:
                fail(f"unsupported xlsx operation: {typ}")
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    inspection = inspect_xlsx_file(out)
    print(json.dumps({"ok": True, "output": str(out), "inspection": inspection}, ensure_ascii=False))


def inspect_csv(args):
    rows = []
    with open(args.input, newline="", encoding=args.encoding) as f:
        reader = csv.reader(f)
        for idx, row in enumerate(reader):
            if idx >= args.max_rows:
                break
            rows.append(row)
    data = {"row_count_sampled": len(rows), "headers": rows[0] if rows else [], "sample_rows": rows[1:]}
    if args.out:
        write_json(args.out, data)
    print(json.dumps({"ok": True, "input": args.input, "inspection": data}, ensure_ascii=False))


def render_office(args):
    soffice = shutil.which("soffice")
    if not soffice:
        fail("soffice is not available in PATH")
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    cmd = [soffice, "--headless", "--convert-to", args.format, "--outdir", str(out_dir), args.input]
    proc = subprocess.run(cmd, text=True, capture_output=True, timeout=args.timeout)
    if proc.returncode != 0:
        fail(f"soffice failed: {proc.stderr or proc.stdout}")
    print(json.dumps({"ok": True, "input": args.input, "out_dir": str(out_dir), "stdout": proc.stdout.strip()}, ensure_ascii=False))


def validate(args):
    path = Path(args.input)
    ext = path.suffix.lower()
    if ext == ".docx":
        inspection = inspect_docx_file(path)
    elif ext == ".xlsx":
        inspection = inspect_xlsx_file(path)
    elif ext == ".csv":
        class A:
            input = str(path)
            encoding = "utf-8"
            max_rows = 20
            out = None
        inspect_csv(A())
        return
    else:
        fail(f"unsupported file type: {ext}")
    expect = read_json(args.expect) if args.expect else {}
    text_blob = json.dumps(inspection, ensure_ascii=False)
    missing = [s for s in expect.get("contains_text", []) if s not in text_blob]
    if missing:
        fail(f"missing expected text: {missing}")
    print(json.dumps({"ok": True, "input": args.input, "inspection": inspection}, ensure_ascii=False))


def self_test(args):
    imports = import_status()
    required = ["docx", "openpyxl", "pandas", "xlsxwriter", "lxml"]
    missing = {k: v for k, v in imports.items() if k in required and v != "ok"}
    if missing:
        fail(f"missing required imports: {missing}")
    work = Path(args.work_dir)
    work.mkdir(parents=True, exist_ok=True)

    docx_spec = {
        "title": "Office Runtime 自检",
        "subtitle": "Word 创建、排版、表格、重读校验",
        "title_align": "center",
        "sections": [{
            "heading": "能力检查",
            "paragraphs": ["这是一份由沙箱内置 office runtime 生成的 Word 文档。"],
            "bullets": [["标题样式", "正文段落", "表格布局"]],
            "tables": [{"headers": ["项目", "状态"], "rows": [["DOCX", "通过"], ["表格", "通过"]]}],
        }],
    }
    docx_spec_path = work / "docx_spec.json"
    write_json(docx_spec_path, docx_spec)
    create_docx(argparse.Namespace(spec=str(docx_spec_path), out=str(work / "selftest.docx")))
    ops_path = work / "docx_ops.json"
    write_json(ops_path, {"operations": [{"type": "replace_text", "old": "通过", "new": "OK"}]})
    edit_docx(argparse.Namespace(input=str(work / "selftest.docx"), ops=str(ops_path), out=str(work / "selftest_edited.docx")))

    xlsx_spec = {
        "sheets": [{
            "name": "Summary",
            "rows": [["项目", "Q1", "Q2"], ["收入", 100, 120], ["成本", 40, 50], ["利润", None, None]],
            "formulas": {"B4": "=B2-B3", "C4": "=C2-C3"},
            "widths": {"A": 16, "B": 12, "C": 12},
            "table": True,
            "charts": [{"type": "bar", "title": "收入趋势", "cats_col": 1, "data_min_col": 2, "data_max_col": 3, "data_max_row": 3, "anchor": "E2"}],
        }]
    }
    xlsx_spec_path = work / "xlsx_spec.json"
    write_json(xlsx_spec_path, xlsx_spec)
    create_xlsx(argparse.Namespace(spec=str(xlsx_spec_path), out=str(work / "selftest.xlsx")))
    xlsx_ops_path = work / "xlsx_ops.json"
    write_json(xlsx_ops_path, {"operations": [{"type": "set_formula", "sheet": "Summary", "cell": "D4", "formula": "=SUM(B4:C4)"}, {"type": "set_width", "sheet": "Summary", "column": "D", "width": 14}]})
    edit_xlsx(argparse.Namespace(input=str(work / "selftest.xlsx"), ops=str(xlsx_ops_path), out=str(work / "selftest_edited.xlsx")))

    result = {"ok": True, "runtime_version": RUNTIME_VERSION, "imports": imports, "soffice": shutil.which("soffice") or "", "work_dir": str(work)}
    if args.require_render and not result["soffice"]:
        fail("soffice is required but not available")
    print(json.dumps(result, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(prog="office")
    sub = parser.add_subparsers(dest="command", required=True)

    sub.add_parser("capabilities").set_defaults(func=capabilities)
    sub.add_parser("fonts").set_defaults(func=list_fonts)
    p = sub.add_parser("self-test")
    p.add_argument("--work-dir", default="office-self-test")
    p.add_argument("--require-render", action="store_true")
    p.set_defaults(func=self_test)

    p = sub.add_parser("create-docx")
    p.add_argument("--spec", required=True)
    p.add_argument("--out", required=True)
    p.set_defaults(func=create_docx)
    p = sub.add_parser("inspect-docx")
    p.add_argument("input")
    p.add_argument("--out")
    p.set_defaults(func=inspect_docx)
    p = sub.add_parser("edit-docx")
    p.add_argument("input")
    p.add_argument("--ops", required=True)
    p.add_argument("--out", required=True)
    p.set_defaults(func=edit_docx)

    p = sub.add_parser("create-xlsx")
    p.add_argument("--spec", required=True)
    p.add_argument("--out", required=True)
    p.set_defaults(func=create_xlsx)
    p = sub.add_parser("inspect-xlsx")
    p.add_argument("input")
    p.add_argument("--out")
    p.set_defaults(func=inspect_xlsx)
    p = sub.add_parser("edit-xlsx")
    p.add_argument("input")
    p.add_argument("--ops", required=True)
    p.add_argument("--out", required=True)
    p.set_defaults(func=edit_xlsx)

    p = sub.add_parser("inspect-csv")
    p.add_argument("input")
    p.add_argument("--out")
    p.add_argument("--encoding", default="utf-8")
    p.add_argument("--max-rows", type=int, default=20)
    p.set_defaults(func=inspect_csv)

    p = sub.add_parser("render")
    p.add_argument("input")
    p.add_argument("--out-dir", required=True)
    p.add_argument("--format", default="pdf")
    p.add_argument("--timeout", type=int, default=120)
    p.set_defaults(func=render_office)

    p = sub.add_parser("validate")
    p.add_argument("input")
    p.add_argument("--expect")
    p.set_defaults(func=validate)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
